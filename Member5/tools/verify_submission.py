"""Read-only integrity/consistency audit of the completed Member 5 package.

Run from any directory with Python, Pillow, PyMuPDF, python-docx and python-pptx.
This does not run the lab, authenticate, or change any evidence.
"""
from pathlib import Path
import ast
import hashlib
import json
import re
import zipfile

import pymupdf as fitz
from PIL import Image, ImageChops
from docx import Document
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'submission'


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def normalize(text):
    return re.sub(r'\s+', '', text.replace('\u00ad', '').replace('`', ''))


def main():
    selected = json.loads((OUT / 'evidence/SELECTION.json').read_text())
    raw = ROOT.parent / 'idk just check/evidence'
    for item in selected:
        for path in (OUT / 'evidence' / item['selected'],
                     ROOT / 'evidence' / item['selected'], raw / item['original']):
            assert sha(path) == item['sha256'], f'Evidence hash mismatch: {path}'
            with Image.open(path) as image:
                image.verify()
    print(f'PASS: {len(selected)} selected PNGs match manifests, both evidence folders and raw captures.')

    crops = json.loads((OUT / 'figures/CROP_MANIFEST.json').read_text())
    for item in crops:
        original = OUT / 'evidence' / item['original']
        figure = OUT / 'figures' / item['figure']
        assert sha(original) == item['original_sha256'], original
        assert sha(figure) == item['figure_sha256'], figure
        with Image.open(original) as image, Image.open(figure) as actual:
            expected = image.crop(item['crop_box_xyxy'])
            assert expected.size == actual.size, figure
            assert ImageChops.difference(expected.convert('RGB'), actual.convert('RGB')).getbbox() is None, figure
    print(f'PASS: {len(crops)} derivatives match recorded hashes and exact original crop pixels.')

    results = sorted((OUT / 'logs').glob('*-result.json'))
    for path in results:
        result = json.loads(path.read_text())
        summary = (OUT / 'logs' / (result['label'] + '-summary.txt')).read_text()
        if result['changed']:
            assert f"CHANGE detected after {result['attempts']} attempts and {result['monitor_elapsed_seconds']} seconds" in summary
            assert result['exact_record_count'] == 1
            assert result['before_sha256'] != result['after_sha256']
        else:
            assert result['before_sha256'] == result['after_sha256']
            assert (f"attempts={result['attempts']}" in summary or
                    f"INTERRUPTED after {result['attempts']} attempts" in summary)
        assert result['runner_ruid'] == result['runner_euid'] == 1000
        assert result['attacker_confirmed_stopped']
        print(f"PASS: {result['label']}: attempts={result['attempts']}, monitor={result['monitor_elapsed_seconds']}s, wall={result['runner_wall_seconds']}s, changed={result['changed']}")

    for path in (OUT / 'automation').glob('*.py'):
        ast.parse(path.read_text(), filename=str(path))
    print('PASS: packaged Python helper syntax.')

    for name in ('REPORT', 'LIVE_DEMO', 'LIVE_COMMANDS', 'SLIDE_NOTES', 'LIVE_PREPARATION'):
        with zipfile.ZipFile(OUT / (name + '.docx')) as archive:
            assert archive.testzip() is None
        document = Document(OUT / (name + '.docx'))
        pdf = fitz.open(OUT / (name + '.pdf'))
        text = normalize(''.join(page.get_text(clip=fitz.Rect(0, 35, page.rect.width, page.rect.height - 35)) for page in pdf))
        missing = [p.text for p in document.paragraphs if p.text.strip() and normalize(p.text) not in text]
        missing += [cell.text for table in document.tables for row in table.rows for cell in row.cells
                    if cell.text.strip() and normalize(cell.text) not in text]
        assert not missing, f'DOCX/PDF text mismatch in {name}: {missing[:3]}'
        markdown = (OUT / (name + '.md')).read_text(encoding='utf-8')
        assert not re.search(r'\[(?:INSERT|OBSERVED|PASS/FAIL)\b', markdown), name
        for target in re.findall(r'!\[[^\]]*\]\(([^)]+)\)', markdown):
            assert (OUT / target).is_file(), target
        assert len(document.inline_shapes) == len(re.findall(r'^!\[', markdown, re.M)), name
        for page in pdf:
            for word in page.get_text('words'):
                assert fitz.Rect(word[:4]) in page.rect, f'Off-page text in {name}, page {page.number + 1}'
        print(f'PASS: {name}: {len(pdf)} PDF pages; {len(document.inline_shapes)} images; DOCX/PDF text matches; no off-page words or result placeholders.')

    deck = Presentation(OUT / 'MAIN_PRESENTATION.pptx')
    pdf = fitz.open(OUT / 'MAIN_PRESENTATION.pdf')
    assert len(deck.slides) == len(pdf)
    for number, slide in enumerate(deck.slides):
        text = normalize(pdf[number].get_text())
        assert slide.notes_slide.notes_text_frame.text.strip(), number + 1
        for shape in slide.shapes:
            if shape.has_text_frame:
                assert normalize(shape.text) in text, f'Slide {number + 1} PPTX/PDF text mismatch: {shape.text}'
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= deck.slide_width + 1000
            assert shape.top + shape.height <= deck.slide_height + 1000
    print(f'PASS: {len(deck.slides)} slides; embedded notes; matching PPTX/PDF text; shapes within slide bounds.')
    print('These checks establish local package consistency, not independent proof of VM execution or oral rehearsal.')


if __name__ == '__main__':
    main()
