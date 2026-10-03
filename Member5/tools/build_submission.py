"""Build the completed Member 5 report, findings deck and live-demo documents.

Run with python-docx, python-pptx and Pillow on PYTHONPATH. Sources are the
Markdown documents in ../submission. Original screenshots are copied verbatim;
every readability crop has an explicit provenance entry.
"""
from pathlib import Path
import argparse
import hashlib
import json
import re
import shutil

from PIL import Image, ImageFont
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pptx import Presentation
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.dml.color import RGBColor as PColor
from pptx.util import Inches as PI, Pt as PP


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'submission'
VM = Path('/home/seed/issd-member5')
AUDIT = VM / 'evidence/audit-20261003-064223'
NAVY = '142A43'
TEAL = '087E8B'
GOLD = 'C48312'
GRAY = '526175'
PALE = 'F2F5F8'
WHITE = 'FFFFFF'
BLACK = '152436'


def digest(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def copy_original(source, dest):
    dest.parent.mkdir(parents=True, exist_ok=True)
    if dest.exists() and digest(source) != digest(dest):
        raise RuntimeError('Existing different artifact: ' + str(dest))
    if not dest.exists():
        shutil.copy2(source, dest)


def assets():
    (OUT / 'figures').mkdir(parents=True, exist_ok=True)
    for item in json.loads((AUDIT / 'selected/SELECTION.json').read_text()):
        source = AUDIT / 'selected' / item['selected']
        if digest(source) != item['sha256']:
            raise RuntimeError('Original evidence hash mismatch')
        copy_original(source, OUT / 'evidence' / source.name)
    copy_original(AUDIT / 'selected/SELECTION.json', OUT / 'evidence/SELECTION.json')
    crops = {
        'report-S01': ('S01-vm-setup.png', (205, 132, 691, 512)),
        'report-S02': ('S02-guest-environment.png', None),
        'report-S03': ('S03-lab-permissions.png', (996, 28, 1920, 552)),
        'report-S04': ('S04-vulnerable-code.png', (72, 28, 996, 610)),
        'report-S05': ('S05-task1-target-validation.png', None),
        'report-S06-A': ('S06-task2a-timing.png', (72, 28, 996, 242)),
        'report-S06-B': ('S06-task2a-timing.png', (996, 28, 1920, 184)),
        'report-S07': ('S07-task2a-result.png', (996, 28, 1920, 543)),
        'report-S08-A': ('S08a-task2b-running.png', (72, 28, 996, 320)),
        'report-S08-B': ('S08b-task2b-running.png', (996, 28, 1920, 502)),
        'report-S09': ('S09-task2b-result.png', (996, 28, 1920, 967)),
        'report-S10-A': ('S10a-task2b-sticky-bit.png', (72, 228, 996, 416)),
        'report-S10-B': ('S10b-task2b-file-exists.png', (996, 535, 1920, 967)),
        'report-S11': ('S11-task2c-atomic-result.png', (996, 28, 1920, 967)),
        'report-S12-trial': ('S12a-task3a-trial.png', (996, 260, 1920, 967)),
        'report-S12-controls': ('S12b-task3a-controls.png', (996, 105, 1920, 967)),
        'report-S13-trial': ('S13a-task3b-trial.png', (996, 28, 1920, 967)),
        'report-S13-controls': ('S13b-task3b-controls.png', (996, 420, 1920, 967)),
        'report-S14': ('S14-cleanup.png', (996, 28, 1920, 967)),
        'slide-record': ('S09-task2b-result.png', (996, 699, 1920, 809)),
        'slide-login': ('S09-task2b-result.png', (996, 583, 1920, 604)),
        'slide-uid': ('S09-task2b-result.png', (996, 906, 1920, 967)),
        'slide-sticky': ('S10b-task2b-file-exists.png', (996, 817, 1920, 967)),
        'slide-denied': ('S13b-task3b-controls.png', (996, 525, 1920, 954)),
    }
    manifest = []
    for name, (original, box) in crops.items():
        source = OUT / 'evidence' / original
        with Image.open(source) as image:
            actual = box or (0, 0, image.width, image.height)
            if not (0 <= actual[0] < actual[2] <= image.width and
                    0 <= actual[1] < actual[3] <= image.height):
                raise RuntimeError('Out-of-bounds crop: ' + name)
            result = image.crop(actual)
            target = OUT / 'figures' / (name + '.png')
            result.save(target)
        manifest.append({'figure': target.name, 'original': original,
                         'original_sha256': digest(source), 'crop_box_xyxy': actual,
                         'purpose': 'Readability excerpt; no alteration of contained pixels',
                         'figure_sha256': digest(target)})
    (OUT / 'figures/CROP_MANIFEST.json').write_text(json.dumps(manifest, indent=2) + '\n')

    for name in ('vulp.c', 'attack_naive.c', 'attack_atomic.c', 'build.sh', 'run_trials.sh', 'input.txt'):
        copy_original(VM / 'lab-files' / name, OUT / 'source' / name)
    attribution = (ROOT / 'README.md').read_text().split('## Sources and attribution', 1)[1]
    (OUT / 'source/SOURCE_ATTRIBUTION.md').write_text(
        '# Original source attribution\n\nRetained from the supplied Member 5 lab pack. '
        'Completed experimental observations are in REPORT.md.\n\n## Sources and attribution'
        + attribution)
    for name in ('reset.sh', 'audit-setup.sh', 'cleanup.sh', 'lab_runner.py', 'verify-record.py',
                 'task3a-controls.sh', 'task3b-controls.sh'):
        copy_original(VM / 'automation' / name, OUT / 'automation' / name)
    for path in (VM / 'lab-files/logs').iterdir():
        if path.is_file() and ('passwd-after' not in path.name) and (
                path.suffix == '.json' or path.name.endswith(('-summary.txt', '-attacker.txt',
                '-runner.txt', '-last-output.txt', '-controls.txt', '-allowed-after.txt'))
                or path.name.startswith('cleanup-') and path.suffix == '.txt'):
            copy_original(path, OUT / 'logs' / path.name)
    for source, name in (
        (AUDIT / 'verified-logins.json', 'audit-verified-logins.json'),
        (VM / 'evidence/automation-20261003-054443/verified-logins.json', 'initial-verified-logins.json'),
        (VM / 'evidence/automation-20261003-054443/exact-append-check.json', 'exact-append-check.json'),
        (VM / 'evidence/automation-20261003-054443/monitor-changes.diff', 'monitor-changes.diff'),
        (AUDIT / 'S08-selected-frame-provenance.json', 'S08-selected-frame-provenance.json'),
        (VM / 'sysctl-before.txt', 'sysctl-before.txt'),
        (VM / 'passwd-before.sha256', 'passwd-before.sha256'),
    ):
        copy_original(source, OUT / 'provenance' / name)
    for source in (AUDIT / 'task2b-audit1-frames').iterdir():
        if source.is_file():
            copy_original(source, OUT / 'provenance/task2b-audit1-frames' / source.name)


def shade(paragraph, fill):
    el = OxmlElement('w:shd')
    el.set(qn('w:fill'), fill)
    paragraph._p.get_or_add_pPr().append(el)


def inline(paragraph, text):
    # Preserve link labels; reference URLs are printed explicitly in References.
    text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1', text)
    for part in re.split(r'(\*\*.*?\*\*|`[^`]+`)', text):
        run = paragraph.add_run(part[2:-2] if part.startswith('**') else
                                part[1:-1] if part.startswith('`') else part)
        if part.startswith('**'):
            run.bold = True
        if part.startswith('`'):
            run.font.name = 'Courier New'
            run.font.size = Pt(9)


def field(paragraph, instruction):
    el = OxmlElement('w:fldSimple')
    el.set(qn('w:instr'), instruction)
    el.set(qn('w:dirty'), 'true')
    run = OxmlElement('w:r')
    properties = OxmlElement('w:rPr')
    size = OxmlElement('w:sz')
    size.set(qn('w:val'), '16')
    properties.append(size)
    run.append(properties)
    text = OxmlElement('w:t')
    text.text = '1'
    run.append(text)
    el.append(run)
    paragraph._p.append(el)


def render_doc(name):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.2677), Inches(11.6929)
    section.top_margin = section.bottom_margin = Inches(.62)
    section.left_margin = section.right_margin = Inches(.65)
    section.header_distance = section.footer_distance = Inches(.27)
    for style_name in ('Normal', 'Body Text', 'List Bullet', 'List Number'):
        style = doc.styles[style_name]
        style.font.name = 'Arial'
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.06
    for level, size in ((1, 17), (2, 13), (3, 11)):
        style = doc.styles['Heading ' + str(level)]
        style.font.name = 'Arial'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(NAVY if level == 1 else TEAL)
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    doc.styles['Title'].font.name = 'Arial'
    doc.styles['Title'].font.size = Pt(29)
    doc.styles['Title'].font.color.rgb = RGBColor.from_string(NAVY)
    if name == 'LIVE_COMMANDS':
        section.top_margin = section.bottom_margin = Inches(.50)
        for style_name in ('Normal', 'Body Text', 'List Bullet', 'List Number'):
            style = doc.styles[style_name]
            style.font.size = Pt(9.5)
            style.paragraph_format.space_after = Pt(3)
            style.paragraph_format.line_spacing = 1
        doc.styles['Title'].font.size = Pt(24)
        for level in (1,2,3):
            style = doc.styles['Heading ' + str(level)]
            style.font.size = Pt(13 if level == 1 else 11)
            style.paragraph_format.space_before = Pt(7)
            style.paragraph_format.space_after = Pt(4)
    header = section.header.paragraphs[0]
    header.text = 'ISSD  /  MEMBER 5                                      SEED RACE-CONDITION LAB'
    for run in header.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(GRAY)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Member 5 — ISSD   |   ')
    field(footer, 'PAGE')
    footer.add_run(' / ')
    field(footer, 'NUMPAGES')
    for run in footer.runs:
        run.font.size = Pt(8)
    doc.core_properties.author = 'Member 5 — ISSD'
    doc.core_properties.title = name.replace('_', ' ').title()
    doc.core_properties.subject = 'Completed findings and live-demo support; actual SEED VM evidence'
    lines = (OUT / (name + '.md')).read_text().splitlines()
    i = 0
    while i < len(lines):
        line = lines[i]
        if line == '<!-- pagebreak -->':
            if name != 'LIVE_DEMO':
                doc.add_page_break()
        elif line.startswith('```'):
            code = []
            i += 1
            while i < len(lines) and not lines[i].startswith('```'):
                code.append(lines[i])
                i += 1
            if i == len(lines):
                raise RuntimeError('Unclosed code block: ' + name)
            p = doc.add_paragraph()
            p.paragraph_format.keep_together = True
            p.paragraph_format.left_indent = Inches(.09)
            p.paragraph_format.right_indent = Inches(.06)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            if name == 'LIVE_COMMANDS':
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1
            shade(p, 'EEF2F6')
            r = p.add_run('\n'.join(code))
            r.font.name = 'Courier New'
            r.font.size = Pt(8.4 if name == 'REPORT' else 9.5)
        elif match := re.match(r'^!\[(.*?)\]\((.*?)\)$', line):
            caption, filename = match.groups()
            image = OUT / filename
            with Image.open(image) as im:
                width = min(6.90, 7.8 * im.width / im.height)
                if image.name == 'report-S14.png':
                    width = min(width, 6.50)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            picture = p.add_run().add_picture(str(image), width=Inches(width))
            picture._inline.docPr.set('descr', caption)
            p = doc.add_paragraph(caption)
            p.paragraph_format.keep_together = True
            p.paragraph_format.space_after = Pt(12)
            for run in p.runs:
                run.italic = True
                run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor.from_string(GRAY)
        elif line.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                row = [x.strip() for x in re.split(r'\|(?=(?:[^`]*`[^`]*`)*[^`]*$)',
                                                  lines[i].strip().strip('|'))]
                if not all(re.fullmatch(r':?-+:?', x) for x in row):
                    rows.append(row)
                i += 1
            table = doc.add_table(rows=0, cols=max(map(len, rows)))
            table.style = 'Light Shading Accent 1'
            for index, row in enumerate(rows):
                cells = table.add_row().cells
                for col, value in enumerate(row):
                    p = cells[col].paragraphs[0]
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.space_before = Pt(3)
                    inline(p, value)
                    for run in p.runs:
                        run.font.size = Pt(9)
                        if index == 0:
                            run.bold = True
                if index == 0:
                    el = OxmlElement('w:tblHeader')
                    table.rows[index]._tr.get_or_add_trPr().append(el)
                el = OxmlElement('w:cantSplit')
                table.rows[index]._tr.get_or_add_trPr().append(el)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue
        elif match := re.match(r'^(#{1,4}) (.*)$', line):
            depth, text = len(match[1]), match[2]
            p = doc.add_paragraph(style='Title') if depth == 1 else doc.add_heading(level=depth - 1)
            inline(p, text)
        elif line.startswith('* '):
            inline(doc.add_paragraph(style='List Bullet'), line[2:])
        elif line.strip() and not line.startswith('<!--'):
            p = doc.add_paragraph()
            if line.startswith('**') and line.endswith('**'):
                p.paragraph_format.keep_with_next = True
            inline(p, line)
        i += 1
    destination = OUT / (name + '.docx')
    doc.save(destination)
    check = Document(destination)
    if len(check.inline_shapes) != len(doc.inline_shapes):
        raise RuntimeError('Image count changed on DOCX reopen')
    print(destination.name, len(check.paragraphs), 'paragraphs;', len(check.inline_shapes), 'images')


class Deck:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = PI(13.3333), PI(7.5)
        self.notes = []
        self.bounds = []

    def text(self, slide, x, y, w, h, text, size=24, color=BLACK,
             bold=False, font='Arial', align=None):
        box = slide.shapes.add_textbox(PI(x), PI(y), PI(w), PI(h))
        box.name = 'Text: ' + text[:45].replace('\n', ' ')
        tf = box.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = 0
        tf.margin_top = tf.margin_bottom = 0
        for i, line in enumerate(text.split('\n')):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.font.name = font
            p.font.size = PP(size)
            p.font.bold = bold
            p.font.color.rgb = PColor.from_string(color)
            p.space_after = PP(6 if size >= 18 else 2)
            p.line_spacing = 1.05
            if align is not None:
                p.alignment = align
        self.bounds.append({'slide': len(self.prs.slides), 'text': text,
                            'box': [x, y, w, h], 'font_pt': size})
        return box

    def rect(self, slide, x, y, w, h, color=PALE, radius=True):
        shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius
                                      else MSO_AUTO_SHAPE_TYPE.RECTANGLE,
                                      PI(x), PI(y), PI(w), PI(h))
        shape.fill.solid()
        shape.fill.fore_color.rgb = PColor.from_string(color)
        shape.line.fill.background()
        if radius:
            shape.adjustments[0] = .08
        return shape

    def base(self, title, subtitle='', tag='FINDINGS', backup=False):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = PColor.from_string('FAFBFD')
        self.rect(slide, 0, 0, 13.3333, .12, TEAL, False)
        self.text(slide, .55, .31, 10, .22, 'ISSD  /  MEMBER 5  /  ' + ('BACKUP — ' if backup else '') + tag,
                  10, TEAL, True)
        self.text(slide, .55, .73, 12.15, .61, title, 33, NAVY, True)
        if subtitle:
            self.text(slide, .57, 1.46, 12.0, .54, subtitle, 17, GRAY)
        self.text(slide, .56, 7.07, 11.5, .18,
                  'SEED Race-Condition Lab • Wenliang Du / SEED Labs • Actual VM findings: 3 Oct 2026', 9, GRAY)
        self.text(slide, 12.20, 7.03, .55, .26, str(len(self.prs.slides)), 12, TEAL, True)
        return slide

    def note(self, slide, heading, text):
        slide.notes_slide.notes_text_frame.text = text
        self.notes.append((len(self.prs.slides), heading, text))

    def picture(self, slide, path, x, y, w, h):
        with Image.open(path) as im:
            scale = min(w / im.width, h / im.height)
            pw, ph = im.width * scale, im.height * scale
        slide.shapes.add_picture(str(path), PI(x + (w-pw)/2), PI(y), PI(pw), PI(ph))


def build_deck():
    d = Deck()
    s = d.prs.slides.add_slide(d.prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = PColor.from_string(NAVY)
    d.rect(s, .63, .62, .15, 5.96, TEAL, False)
    d.text(s, 1.00, .67, 10, .35, 'MEMBER 5 — ISSD', 15, '73D8D2', True)
    d.text(s, 1.00, 1.42, 11.2, 1.55, 'One name.\nTwo different files.', 47, WHITE, True)
    d.text(s, 1.02, 3.37, 10.9, .66, 'Race-condition vulnerability lab', 28, WHITE)
    d.text(s, 1.03, 4.27, 10.8, .8,
           'A race condition is timing-dependent. Here the target changes between check and use.', 24, 'D9E5F0')
    d.rect(s, 1.02, 5.72, 10.95, .66, '203E5E')
    d.text(s, 1.25, 5.91, 10.45, .31, 'MECHANISM  →  LIVE TERMINALS  →  FINDINGS  →  DEFENCES', 17, '8DE4DF', True)
    d.note(s, 'Introduce the race before the demonstration',
           'Introduce Member 5 and the SEED Race-Condition Lab. A race condition means the outcome depends on the ordering of concurrent operations. Our case is time of check to time of use: a program checks one target and later opens another target through the same name. This is a local application-level lab, not Dirty COW or a remote attack. The actual terminal demonstration comes after the mechanism and the audience cues.')

    s = d.base('Where the race lives', 'The pathname stays the same; the object reached through it can change.', 'MECHANISM')
    d.rect(s, .57, 2.20, 7.28, 3.87, NAVY)
    d.text(s, .90, 2.51, 6.6, .35, 'CHECK', 14, '7DDDD5', True)
    d.text(s, .90, 3.01, 6.57, .62, 'access("/tmp/XYZ", W_OK)', 25, WHITE, font='Courier New')
    d.text(s, .92, 3.77, 6.3, .38, 'Attacker switches the symlink here', 20, 'F6CE84', True)
    d.text(s, .90, 4.43, 6.6, .35, 'USE', 14, '7DDDD5', True)
    d.text(s, .90, 4.92, 6.57, .6, 'fopen("/tmp/XYZ", "a+")', 25, WHITE, font='Courier New')
    d.rect(s, 8.20, 2.20, 4.52, 1.75)
    d.text(s, 8.45, 2.48, 4.0, .38, 'Real UID: 1000 (seed)', 23, TEAL, True)
    d.text(s, 8.45, 3.04, 3.95, .65, 'The permission check uses the ordinary caller.', 20)
    d.rect(s, 8.20, 4.23, 4.52, 1.84)
    d.text(s, 8.45, 4.51, 4.0, .38, 'Effective UID: 0', 23, GOLD, True)
    d.text(s, 8.45, 5.07, 3.95, .68, 'The Set-UID victim supplies root authority to the open.', 20)
    d.text(s, .65, 6.42, 12.0, .36, 'Lab prerequisite: deliberately installed root-owned Set-UID victim; attack controls set to 0/0.', 16, GRAY)
    d.note(s, 'Explain real versus effective identity',
           'Point to the two actual operations. access uses the real user identity. fopen resolves the name again and uses the process credentials for the open. In our deliberately configured Set-UID victim the effective identity is root. The attacker changes the symlink from writable /dev/null to protected /etc/passwd in the interval. A symlink does not grant privileges by itself. Do not imply that every Ubuntu installation has this program or that file permissions disappeared. Source: vulp.c, S03/S04, and the SEED task PDF.')

    s = d.base('What to watch in the live demo', 'Switch to the two terminals now. Explain the output as it appears.', 'LIVE DEMO 1')
    d.rect(s, .57, 2.21, 5.98, 2.95, NAVY)
    d.text(s, .88, 2.53, 5.33, .50, 'A — Attacker', 29, WHITE, True)
    d.text(s, .88, 3.28, 5.22, 1.30, 'First: XYZ → /dev/null\nThen: XYZ → /etc/passwd', 25, 'DDE9F4')
    d.rect(s, 6.79, 2.21, 5.95, 2.95)
    d.text(s, 7.10, 2.53, 5.22, .50, 'B — Victim and results', 28, TEAL, True)
    d.text(s, 7.10, 3.28, 5.13, 1.39, 'Check passed\nInspect the complete record\nNon-sudo login → id', 24)
    d.rect(s, .60, 5.55, 12.12, .83, 'FFF0D7')
    d.text(s, .86, 5.76, 11.52, .38, 'Task 2.A: an explicit artificial 10-second delay for explanation.', 23, '815300', True)
    d.note(s, 'Switch to LIVE_DEMO Part 1',
           'Use LIVE_DEMO Part 0 before class, then Part 1 in its exact A/B order. A initializes the /dev/null link. B runs vulp_slow. Only after Check passed, switch A to /etc/passwd during the ten-second wait. Wait for the victim to return, inspect the full record, then run non-sudo su - test and id. Password input is interactive. Explain that this is the official slow-machine simulation, not the no-delay result. Exit the root shell and reset before the next ordinary-user trial. Return to slide 4. If timing misses, say so and use actual saved S06/S07 as recorded evidence.')

    s = d.base('Read the output in the right order', 'The proof is the full chain, not a successful-looking message.', 'INTERPRETATION')
    entries = [
        ('uid=1000(seed)', 'The attack starts as an ordinary user.'),
        ('Exact seven-field test record', 'The chosen account entry reached /etc/passwd.'),
        ('su - test  →  id: uid=0(root)', 'That non-sudo login actually obtained root authority.'),
    ]
    for i, (left, right) in enumerate(entries):
        y = 2.17 + i * 1.08
        d.rect(s, .62, y, 12.05, .91)
        d.text(s, .85, y+.18, 5.55, .55, left, 23, TEAL, True)
        d.text(s, 6.75, y+.18, 5.54, .58, right, 20)
    d.rect(s, .64, 5.85, 12.00, .70, NAVY)
    d.text(s, .88, 6.04, 11.55, .34, 'A changed SHA-256 is a signal to inspect — not proof of root access.', 22, WHITE, True)
    d.note(s, 'Explain why the observed output proves the gain',
           'The first identity output shows seed. The exact-record check establishes that the complete 43-character entry, with seven fields and UID/GID zero, exists once. The meaningful final proof is an actual non-sudo su login followed by id. whoami may say root because both account names map to UID zero. A hash change, a program exit status zero, or an account name alone is insufficient. In Task 1 we inserted the record administratively only to validate it; in the actual attack the privileged victim performed the append. Our saved results S07, S09 and S11 preserve that distinction.')

    s = d.base('The no-delay experiment also succeeded', 'Recorded measurements from 3 October 2026 — a live run can differ.', 'LIVE DEMO 2 / FINDINGS')
    xs = [.82, 5.02, 7.30, 10.07]
    for x, text, width in zip(xs, ['Recorded run', 'Attempts', 'Wall time', 'Verified'], [4.0, 2.0, 2.4, 2.0]):
        d.text(s, x, 2.17, width, .4, text, 20, GRAY, True)
    rows = [('Naive: first success', '20', '0.865 s', 'UID 0'),
            ('Atomic exchange', '1', '0.273 s', 'UID 0'),
            ('Naive: recorded repeat', '1', '0.238 s', 'UID 0')]
    for i, row in enumerate(rows):
        y = 2.77 + i * .78
        d.rect(s, .61, y-.10, 12.05, .69, PALE)
        for x, value, width in zip(xs, row, [4.08, 1.90, 2.4, 2.0]):
            d.text(s, x, y+.04, width, .43, value, 23, TEAL if x == xs[-1] else BLACK, x == xs[-1])
    d.text(s, .73, 5.40, 11.89, .56, 'Same no-delay vulp. Bash timer: whole seconds; a 0 s reading does not mean zero runtime.', 17, GRAY)
    d.text(s, .73, 6.15, 11.89, .42, 'Live demonstration: original vulp + atomic switching, with no artificial delay.', 20, TEAL, True)
    d.note(s, 'Switch to the genuine no-delay atomic demonstration',
           'These are historical measurements, rounded to three decimals on the slide; the report contains exact values and labels. They do not promise that the next trial will finish in one attempt or that one method is always faster. The supplied C victim has no artificial sleep. If presenting live, use Part 2: B resets and creates a unique shared label, A starts the bounded atomic attacker, and B starts the 30-second no-delay monitor only after initialization. Inspect any changed file and verify non-sudo login/id; then exit root and reset. If the 30 seconds expire unchanged, state that the live attempt did not win and use the labelled actual S09/S11 evidence. Do not call that short class attempt the original 300-second experiment. Return to slide 6.')

    s = d.base('Why the naive attacker got stuck', 'File exists and Operation not permitted explained a real failure.', 'SURPRISING OBSERVATION')
    labels = [('1', 'unlink()', 'XYZ is temporarily absent.'),
              ('2', 'fopen("a+")', 'The victim can create a root-owned regular file.'),
              ('3', 'Sticky /tmp', 'seed cannot unlink that root-owned entry.')]
    for i, (number, title, body) in enumerate(labels):
        x = .62 + i * 4.17
        d.rect(s, x, 2.25, 3.93, 2.65)
        d.text(s, x+.24, 2.48, 3.41, .4, number + '  ' + title, 24, TEAL, True)
        d.text(s, x+.24, 3.23, 3.37, 1.26, body, 23)
    d.text(s, .78, 5.25, 11.79, .45, 'Observed: root:seed, mode 0664, under root-owned mode-1777 /tmp.', 20, GRAY)
    d.rect(s, .67, 5.97, 12.0, .61, NAVY)
    d.text(s, .92, 6.12, 11.46, .31, 'RENAME_EXCHANGE removes this missing-name gap — it does not fix the victim.', 20, WHITE, True)
    d.note(s, 'Explain the attacker’s own race and its repair',
           'The unlink/create sequence is not atomic. If the victim already passed access and reaches fopen while the name is absent, a+ can create a root-owned regular file. The attacker can then encounter File exists and later an unlink permission error. The sticky directory rule concerns deletion of an entry; group write permission on the file does not grant that deletion. In the old partial log, repeated successful victim exits did not prove an append to /etc/passwd. We preserved the file, metadata and logs before administrator-assisted reset. The improved attacker creates two links first and atomically exchanges them with RENAME_EXCHANGE; initialize fully before the victim starts. This fixes the attacker’s gap, while the victim still checks and opens separately.')

    s = d.base('Two defences, tested independently', 'The long trials are recorded findings; the live check is a short control.', 'LIVE DEMO 3 / DEFENCES')
    for x, title, condition, metric, body in [
        (.62, 'Least privilege', 'vulp_least • controls 0/0', '9,455 attempts / 300 s',
         'Target hash unchanged\nProtected operation denied\nOrdinary-file write succeeded'),
        (6.84, 'Symlink protection', 'original vulp • controls 1/0', '9,177 attempts / 300 s',
         'Target hash unchanged\nStable privileged open denied\nOwnership/context rule applies')]:
        d.rect(s, x, 2.18, 5.90, 3.71)
        d.text(s, x+.26, 2.47, 5.37, .44, title, 27, TEAL, True)
        d.text(s, x+.26, 3.06, 5.36, .38, condition, 19, GRAY)
        d.text(s, x+.26, 3.65, 5.37, .44, metric, 26, NAVY, True)
        d.text(s, x+.26, 4.35, 5.25, 1.31, body, 21)
    d.text(s, .80, 6.25, 11.85, .47, 'Observed: no protected-file change in either 300-second trial.', 21, TEAL, True)
    d.note(s, 'Explain both controls; optionally show the stable-link denial live',
           'For the first experiment we changed the program and left both OS controls disabled. seteuid(getuid()) drops effective privilege before access and fopen and keeps it dropped through writing/closing. The allowed-file control confirmed normal functionality. For the second we returned to the original vulnerable victim and enabled only protected_symlinks, leaving protected_regular at zero. The kernel blocks a root-effective follower of a seed-owned link under sticky root-owned /tmp in this case. In Part 3 of the live guide, a stable /dev/null link should produce Open failed: Permission denied even though the real-user check can pass. That is a short explanatory control, not a live repetition of five minutes. Finite negative trials alone are not universal proof. These two mechanisms address different layers.')

    s = d.base('Protect the operation, not just the filename', 'A local file-handling flaw can undermine the account database used by authentication.', 'TAKEAWAY')
    cards = [('Use the right privilege', 'Drop privilege before user-directed file operations.'),
             ('Use the opened object', 'Open once; retain and validate the file descriptor.'),
             ('Create files safely', 'Use mkstemp, O_CREAT|O_EXCL and trusted directories.')]
    for i, (title, body) in enumerate(cards):
        x = .64 + i * 4.16
        d.rect(s, x, 2.29, 3.91, 2.47)
        d.text(s, x+.23, 2.57, 3.41, .87, title, 25, TEAL, True)
        d.text(s, x+.23, 3.62, 3.40, 1.08, body, 20)
    d.text(s, .79, 5.16, 11.8, .72, 'Do not check an attacker-changeable name and later use it with more privilege.', 28, NAVY, True)
    d.text(s, .81, 6.19, 11.7, .52, 'After the demo: exit root, stop processes, restore the baseline and remove Set-UID (Part 4).', 18, GRAY)
    d.note(s, 'Close the main presentation and perform cleanup',
           'The gain was a verified local UID-zero login, not merely a changed text file. Correct account-database handling supports authentication integrity. Emphasize secure descriptor-based handling, appropriate permissions/trusted directories, safe temporary creation and only the required privilege. A security-relevant atomic operation is useful; merely making the attacker faster does not secure the victim. OS symlink protection is defence in depth and context-specific. Use Part 4 after class or rehearsal: the supplied cleanup restores the original password file, removes the test paths and Set-UID, and applies the explicitly chosen 1/2 runtime policy. The saved 0/0 file is not asserted to be original defaults. Stop the main talk here; the remaining slides are supporting material.')

    s = d.base('Recorded verification: account and UID', 'Actual image excerpts; whole originals are retained as S09.', 'EVIDENCE', True)
    d.text(s, .75, 2.02, 11.9, .35, 'Complete-record check', 21, TEAL, True)
    d.picture(s, OUT/'figures/slide-record.png', .76, 2.48, 11.74, 1.85)
    d.text(s, .75, 4.42, 11.9, .35, 'Actual non-sudo command and UID output (separate excerpts)', 21, TEAL, True)
    d.picture(s, OUT/'figures/slide-login.png', .76, 4.95, 11.74, .40)
    d.picture(s, OUT/'figures/slide-uid.png', .76, 5.58, 11.74, .97)
    d.note(s, 'Evidence backup',
           'Use only as recorded evidence, not as a substitute disguised as a live result. These are labelled readability crops from S09, task2b-audit1-20261003-064223. The complete original screenshot and full terminal/record checks are in the evidence package. The current live run may have different output or counts. The fragments are not a newly manufactured terminal screen.')

    s = d.base('What the failures and timers mean', 'Preserved outcomes matter as much as the successful screenshots.', 'RESULT DETAILS', True)
    lines = [('Original naive log', 'Last progress 76,000 / 227 s; final totals unknown.'),
             ('Naive retry 1', '9 attempts; 0 integer seconds; sticky-file failure.'),
             ('Supplemental capture run', '2 attempts; 0 integer seconds; sticky-file failure.'),
             ('Integer timer', '0 seconds means below the timer’s granularity.'),
             ('Runner wall time', 'Includes monitor startup/exit; not the race-window size.')]
    for i, (label, text) in enumerate(lines):
        y = 2.11+i*.81
        d.rect(s, .64, y, 12.0, .68)
        d.text(s, .87, y+.14, 4.00, .42, label, 21, TEAL, True)
        d.text(s, 5.11, y+.14, 7.25, .45, text, 19)
    d.note(s, 'Measurement and failure Q&A',
           'The original log stopped without a final summary, so we do not invent its final totals. The old XYZ file and original logs were preserved before reset. Hash change requires record/login verification, and exit status zero can just mean an append to /dev/null or the wrong temporary file. The wall times include orchestration overhead; their small values do not establish a general success probability or throughput comparison. Monitor and screenshot activity affects scheduling.')

    s = d.base('Live-demo recovery', 'Use the real output; keep the experiment boundary clear.', 'PRESENTER BACKUP', True)
    recover = [('Timing miss / 30 s unchanged', 'State the failed bounded attempt. Use the labelled recorded S09/S11 result.'),
               ('File exists / unlink denied', 'Stop first. Inspect /tmp and XYZ ownership. Preserve the failure; reset in B.'),
               ('Root shell still open', 'exit once, then id. Return to seed before reset or a new ordinary-user trial.'),
               ('Set-UID missing', 'Use Part 0 setup. Do not run the attacker or victim through sudo.')]
    for i, (title, body) in enumerate(recover):
        y = 2.08+i*1.02
        d.text(s, .77, y, 4.22, .70, title, 22, TEAL, True)
        d.text(s, 5.30, y, 7.08, .75, body, 21)
    d.note(s, 'Recovery instructions',
           'Keep LIVE_DEMO.pdf available off the projected terminals. Follow its reset and cleanup sequence. A says Attacker, B says Victim and results; both start as seed. The same run label is passed via class-demo-label.txt. Attack initialization must appear before starting the victim. The supplied runner is bounded and stops the attacker when the monitor ends. If any setup helper reports an error, address that actual error before continuing. Authentication is interactive; no password is embedded in any script.')

    s = d.base('Sources and supporting files', 'Original lab content and actual experimental evidence are distinguishable.', 'REFERENCES', True)
    d.text(s, .81, 2.21, 11.8, 2.61,
           'Wenliang Du / SEED Labs — Race Condition Vulnerability Lab (2006–2020)\n'
           'seedsecuritylabs.org/Labs_20.04/Software/Race_Condition/\n'
           'Linux kernel filesystem sysctl documentation; access(2), rename(2), open(2)\n'
           'Supplied Member 5 allocation and our VM logs / S01–S15 evidence', 22, NAVY)
    d.text(s, .81, 5.21, 11.8, .65,
           'REPORT.pdf: analysis and full results     LIVE_DEMO.pdf: exact A/B command order', 20, TEAL, True)
    d.text(s, .81, 6.06, 11.8, .55,
           'CC BY-NC-SA 4.0 attribution retained. AI-assisted orchestration/document preparation is disclosed in the report.', 16, GRAY)
    d.note(s, 'Attribution and completion status',
           'The teaching content and vulnerable-program pattern are attributed to Wenliang Du / SEED Labs under CC BY-NC-SA 4.0. See the report for full references, source modifications and assistance disclosure. Actual commands/evidence were collected in the user’s VM; authentication input remained interactive. The practical and documents are complete, but a future oral presentation/rehearsal or submission upload is not claimed to have already occurred. Unknown personal/admin fields were omitted at the user’s request.')

    d.prs.core_properties.title = 'Race condition — findings and live demonstration'
    d.prs.core_properties.author = 'Member 5 — ISSD'
    d.prs.core_properties.subject = 'Findings interpretation with terminal-demo transitions at slides 3, 5 and 7'
    target = OUT / 'MAIN_PRESENTATION.pptx'
    d.prs.save(target)
    reopened = Presentation(target)
    if len(reopened.slides) != len(d.notes):
        raise RuntimeError('Missing slides/notes on PPTX reopen')
    for item in d.bounds:
        x,y,w,h = item['box']
        if x < 0 or y < 0 or x+w > 13.34 or y+h > 7.51:
            raise RuntimeError('Off-slide text box: ' + str(item))
    (OUT/'provenance/SLIDE_LAYOUT.json').write_text(json.dumps(d.bounds, indent=2)+'\n')
    notes = ['# Slide notes — Member 5 — ISSD', '',
             'The main talk is slides 1–8; slides 9–12 are supporting evidence/Q&A. No fixed speaking length is imposed. Use LIVE_DEMO.pdf for exact terminal commands.', '',
             'Live transitions: **slide 3 → Part 1**, **slide 5 → Part 2**, **slide 7 → Part 3**. Finish with **Part 4 cleanup**. Recorded findings are not predetermined live output.', '']
    for number, heading, text in d.notes:
        notes += ['## Slide {} — {}'.format(number, heading), '', text, '']
    (OUT/'SLIDE_NOTES.md').write_text('\n'.join(notes))
    print(target.name, len(reopened.slides), 'slides with embedded notes')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--only', choices=('all','assets','docs','slides'), default='all')
    args = parser.parse_args()
    if args.only in ('all','assets'):
        assets()
    if args.only in ('all','slides'):
        build_deck()
    if args.only in ('all','docs'):
        for name in ('REPORT','LIVE_DEMO','LIVE_COMMANDS','SLIDE_NOTES','LIVE_PREPARATION'):
            render_doc(name)


if __name__ == '__main__':
    main()
