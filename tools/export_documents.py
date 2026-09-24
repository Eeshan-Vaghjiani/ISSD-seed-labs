"""Export this pack's Markdown documents to readable Word documents.

Requires: python-docx (pip install python-docx).
Run from any directory: python path/to/Member5/tools/export_documents.py
Generated documents remain templates until actual evidence is inserted.
"""
from pathlib import Path
import re

from docx import Document
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCUMENTS = (
    "START_TO_FINISH_GUIDE",
    "REPORT_TEMPLATE",
    "PRESENTATION_PLAN",
)


def inline(paragraph, text):
    """Render basic emphasis/code and retain URLs in printable text."""
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    for part in re.split(r"(\*\*.*?\*\*|`[^`]+`)", text):
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
        else:
            paragraph.add_run(part)


def export(name):
    source = ROOT / f"{name}.md"
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.7)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    for level in range(1, 4):
        document.styles[f"Heading {level}"].font.color.rgb = RGBColor.from_string("17365D")
    section.footer.paragraphs[0].text = "ISSD | Member 5 | SEED Race-Condition Lab"
    document.core_properties.title = name.replace("_", " ")
    document.core_properties.subject = "Guide / template; insert actual lab evidence"
    lines = source.read_text(encoding="utf-8").splitlines()
    index = 0
    code = False
    while index < len(lines):
        line = lines[index]
        if line.startswith("```"):
            code = not code
            index += 1
            continue
        if code:
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1
            run = paragraph.add_run(line)
            run.font.name = "Consolas"
            run.font.size = Pt(8)
        elif line.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].startswith("|"):
                # A pipe inside an inline-code span is content, not a cell edge.
                row = [cell.strip() for cell in re.split(
                    r"\|(?=(?:[^`]*`[^`]*`)*[^`]*$)",
                    lines[index].strip().strip("|"),
                )]
                if not all(re.fullmatch(r":?-+:?", cell) for cell in row):
                    rows.append(row)
                index += 1
            table = document.add_table(rows=0, cols=max(map(len, rows)))
            table.style = "Light Shading Accent 1"
            for row_number, row in enumerate(rows):
                cells = table.add_row().cells
                for column, value in enumerate(row):
                    inline(cells[column].paragraphs[0], value)
                    if row_number == 0:
                        for run in cells[column].paragraphs[0].runs:
                            run.bold = True
            document.add_paragraph()
            continue
        elif match := re.match(r"^(#{1,4})\s+(.*)", line):
            inline(document.add_heading(level=len(match[1])), match[2])
        elif line.startswith("* "):
            inline(document.add_paragraph(style="List Bullet"), line[2:])
        elif re.match(r"^\d+\. ", line):
            # Keep explicit numbering, including discontinuous procedural steps.
            inline(document.add_paragraph(), line)
        elif line.strip():
            inline(document.add_paragraph(), line)
        index += 1
    destination = ROOT / f"{name}.docx"
    document.save(destination)
    reopened = Document(destination)
    assert len(reopened.paragraphs) == len(document.paragraphs)
    assert len(reopened.tables) == len(document.tables)
    assert not code, f"Unclosed code fence in {source}"
    print(f"Created and reopened {destination} "
          f"({len(reopened.paragraphs)} paragraphs, {len(reopened.tables)} tables)")


if __name__ == "__main__":
    for document_name in DOCUMENTS:
        export(document_name)
