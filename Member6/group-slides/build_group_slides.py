"""Build Member 6's four group-presentation slides (Dirty COW lab).

Reads the shared group deck, clones Member 5's TITLE_ONLY style (Russo One /
Roboto, navy cards 161C44, cyan 29C7FE / pink ED218F, code block 0B0F26),
appends the four slides, and saves:
  * "ISSD presentation (Member 6 added).pptx"  - full deck, Member 6 inserted
    after slide 31, leftover template slides (32-51) removed.
  * "Member 6 slides only.pptx" - just the four slides, same theme.

Run:  python Member6/group-slides/build_group_slides.py
"""
from pathlib import Path
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path("/home/DHB/Downloads/ISSP/ISSD-seed-labs")
SRC = ROOT / "ISSD presentation.pptx"
OUT_MERGED = ROOT / "ISSD presentation (Member 6 added).pptx"
OUT_DIR = ROOT / "Member6/group-slides"
OUT_SINGLE = OUT_DIR / "Member 6 slides only.pptx"
ASSETS = OUT_DIR / "assets"
S07 = ROOT / "Member6/submission/evidence/M6-S07-task1-result.png"
S10 = ROOT / "Member6/submission/evidence/M6-S10-task2-root-proof.png"

NAVY = RGBColor(0x16, 0x1C, 0x44)
CODE = RGBColor(0x0B, 0x0F, 0x26)
CYAN = RGBColor(0x29, 0xC7, 0xFE)
PINK = RGBColor(0xED, 0x21, 0x8F)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0xC9, 0xD3, 0xE3)
CODETXT = RGBColor(0xE6, 0xED, 0xF7)

# Tight strips cropped from the genuine 9 Oct 2026 framebuffer captures.
STRIPS = {
    "task1-result.png": (S07, (66, 158, 1240, 328)),   # diff + cat + ls + stat lines
    "task2-uid0.png": (S10, (66, 230, 670, 348)),      # id / uid=0(root) / id -u 0 / whoami root
}


def crop_strips():
    ASSETS.mkdir(parents=True, exist_ok=True)
    sizes = {}
    for name, (source, box) in STRIPS.items():
        target = ASSETS / name
        with Image.open(source) as image:
            image.crop(box).save(target)
            sizes[name] = image.crop(box).size
    return sizes


def txt(slide, x, y, w, h, text, size, color, bold=False, font="Roboto",
        align=PP_ALIGN.LEFT, wrap=True):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.word_wrap = wrap
    for margin in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(frame, margin, 0)
    for index, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        run = paragraph.add_run()
        run.text = line
        run.font.name = font
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
    return box


def card(slide, x, y, w, h, fill=NAVY):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    shape.shadow.inherit = False
    return shape


def marker(slide, x, y, color):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x), Inches(y), Inches(x + 0.26), Inches(y))
    line.line.color.rgb = color
    line.line.width = Pt(1.75)
    line.shadow.inherit = False


def title(slide, text):
    holder = slide.shapes.title
    if holder is None:  # fall back to a plain styled box
        holder = txt(slide, 0.79, 0.42, 8.42, 0.75, text, 26, WHITE, font="Russo One", align=PP_ALIGN.CENTER)
        return
    frame = holder.text_frame
    frame.text = text
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.CENTER
    for run in paragraph.runs:
        run.font.name = "Russo One"
        run.font.size = Pt(26)
        run.font.color.rgb = WHITE


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def slide_concept(slide):
    title(slide, "Dirty COW: the kernel's copy-on-write race")
    txt(slide, 0.79, 1.30, 8.42, 0.45,
        "A read-only private map must never change the file behind it. A kernel race broke that rule.",
        18, WHITE)
    card(slide, 0.79, 1.95, 5.5, 2.3, CODE)
    txt(slide, 1.05, 2.20, 5.1, 1.9,
        'fd = open("/zzz", O_RDONLY)   /* no write right */\n'
        "map = mmap(fd, PROT_READ, MAP_PRIVATE)\n"
        'pwrite("/proc/self/mem", data, addr)\n'
        "madvise(map, MADV_DONTNEED)   /* repeat */\n"
        "/* the kernel loses the private copy */",
        16, CODETXT)
    marker(slide, 6.29, 2.34, CYAN)
    txt(slide, 6.60, 2.22, 3.0, 0.3, "We run as seed (UID 1000)", 14, CYAN)
    marker(slide, 6.29, 3.45, PINK)
    txt(slide, 6.60, 3.33, 3.0, 0.3, "Kernel bug: file bytes change", 14, PINK)
    txt(slide, 0.79, 4.50, 8.42, 0.7,
        "Nothing here opens the file for writing - the race is inside the kernel's copy-on-write code (CVE-2016-5195).",
        16, WHITE)
    notes(slide, "Member 6 - Dirty COW (CVE-2016-5195).\n"
                 "Member 5's race was inside a root-owned program; this one is inside the kernel itself.\n"
                 "Walk through the four lines, then switch to the live demo.")


def slide_demo(slide):
    title(slide, "Live demonstration")
    txt(slide, 0.79, 1.30, 8.42, 0.45,
        "/zzz belongs to root, mode 0644. Seed can read it - but never write it.", 18, WHITE)
    cards = (
        (0.79, CYAN, "1", "Show the file is protected",
         "cat /zzz shows the original. echo > /zzz: Permission denied."),
        (3.65, CYAN, "2", "Normal copy-on-write",
         "./cow_control changes its private copy only - the file stays the same."),
        (6.51, PINK, "3", "Run the race",
         "bash run_trial.sh dummy 30 demo - it stops itself when /zzz changes."),
    )
    for x, color, number, heading, body in cards:
        card(slide, x, 2.0, 2.7, 2.95)
        txt(slide, x + 0.22, 2.2, 2.3, 0.6, number, 34, color, font="Roboto")
        txt(slide, x + 0.22, 2.9, 2.3, 0.8, heading, 16, WHITE, bold=True)
        txt(slide, x + 0.22, 3.8, 2.3, 1.0, body, 14, MUTED)
    txt(slide, 0.79, 5.05, 8.42, 0.4,
        "If the race misses this time, the next slide shows the recorded 9 October run.", 14, MUTED)
    notes(slide, "Live demo in the SEED VM as seed (VM: ISSD-Member6-SEED12).\n"
                 "Prep script already created /zzz before class - see Member6/live-demo/DEMO_SCRIPT.md.\n"
                 "Type the six commands slowly; every one is on the demo cheat-sheet.")


def slide_result(slide, strip_sizes):
    title(slide, "Result: file changed, then uid 0")
    txt(slide, 0.79, 1.14, 8.42, 0.32,
        "Task 1 - seed rewrote the root-owned file /zzz (recorded)", 14, CYAN, bold=True)
    width, height = strip_sizes["task1-result.png"]
    h = 8.42 * height / width
    slide.shapes.add_picture(str(ASSETS / "task1-result.png"), Inches(0.79), Inches(1.52), width=Inches(8.42))
    txt(slide, 0.79, 2.86, 8.42, 0.32,
        "Task 2 - charlie's UID 1001 -> 0000; fresh login without sudo gives uid 0", 14, CYAN, bold=True)
    width, height = strip_sizes["task2-uid0.png"]
    slide.shapes.add_picture(str(ASSETS / "task2-uid0.png"), Inches(0.79), Inches(3.24), width=Inches(8.42))
    txt(slide, 0.79, 4.92, 8.42, 0.55,
        "Recorded 9 Oct 2026 - one bounded 30 s run each, 0 s elapsed.\n"
        "First su attempt failed, retry succeeded. After the demo the account was restored to UID 1001.",
        13, MUTED)
    notes(slide, "Both are real runs from 9 October (evidence retained with hashes).\n"
                 "Task 1: whole file became 111111******333333, root:root 0644 kept.\n"
                 "Task 2: only the four UID digits changed; id -u printed 0, whoami root.")


def slide_defence(slide):
    title(slide, "Countermeasures")
    cards = (
        (0.79, "PATCH THE KERNEL", "Fixed 2016", "Upstream fix landed in October 2016.",
         "Supported distros ship it. Our VM keeps the old kernel to demo the bug."),
        (3.65, "STAY UPDATED", "Update", "A fix only helps when the fixed kernel is actually running.",
         "Update and reboot. Retire end-of-life releases like this 12.04 VM."),
        (6.51, "LIMIT & MONITOR", "Watch uid 0", "Fewer local users, fewer chances to race.",
         "Baseline /etc/passwd; alert on new UID-0 entries and strange logins."),
    )
    for x, header, big, sub, body in cards:
        card(slide, x, 1.18, 2.7, 2.6)
        txt(slide, x + 0.22, 1.38, 2.3, 0.35, header, 13, CYAN, bold=True)
        txt(slide, x + 0.22, 1.73, 2.3, 0.5, big, 24, WHITE, font="Russo One")
        txt(slide, x + 0.22, 2.28, 2.3, 0.5, sub, 13, MUTED)
        txt(slide, x + 0.22, 2.86, 2.3, 0.85, body, 14, WHITE)
    card(slide, 0.79, 5.09, 0.08, 0.33, PINK)
    txt(slide, 1.10, 5.09, 8.0, 0.35,
        "Private memory must stay private - the kernel fix is the real defence.", 16, WHITE, bold=True)
    notes(slide, "Four defences: patch, update, limit, monitor.\n"
                 "Deleting the spare template slides was this member's job as the last presenter.")


def build(prs, layout, strip_sizes):
    for builder in (slide_concept, slide_demo,
                    lambda s: slide_result(s, strip_sizes), slide_defence):
        slide = prs.slides.add_slide(layout)
        builder(slide)


def drop(ids_to_drop, prs):
    lst = prs.slides._sldIdLst
    for element in ids_to_drop:
        prs.part.drop_rel(element.get(qn("r:id")))
        lst.remove(element)


def dedupe(path):
    """Rewrite the zip keeping one entry per name (Google exports can yield
    duplicate partnames whose contents are identical)."""
    import zipfile
    source = zipfile.ZipFile(path)
    target = path.with_suffix(".tmp.pptx")
    with zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as output:
        for name in dict.fromkeys(source.namelist()):
            output.writestr(name, source.read(name))
    source.close()
    target.replace(path)


def main():
    strip_sizes = crop_strips()
    layout = None

    prs = Presentation(SRC)
    assert len(prs.slides) == 51
    layout = prs.slides[27].slide_layout  # Member 5's TITLE_ONLY
    before = list(prs.slides._sldIdLst)
    build(prs, layout, strip_sizes)
    drop(before[31:51], prs)  # leftover template slides
    prs.save(OUT_MERGED)
    dedupe(OUT_MERGED)

    single = Presentation(SRC)
    before = list(single.slides._sldIdLst)
    build(single, layout, strip_sizes)
    drop(before, single)
    single.save(OUT_SINGLE)
    dedupe(OUT_SINGLE)

    merged = Presentation(OUT_MERGED)
    assert len(merged.slides) == 35, len(merged.slides)
    assert merged.slides[31].shapes.title.text.startswith("Dirty COW")
    assert len(Presentation(OUT_SINGLE).slides) == 4
    print("Slides built:", len(merged.slides), "merged /", len(Presentation(OUT_SINGLE).slides), "standalone")
    print("Strips:", {k: f"{w}x{h}" for k, (w, h) in strip_sizes.items()})
    print("Saved:", OUT_MERGED.name, "|", OUT_SINGLE.name)


if __name__ == "__main__":
    main()
