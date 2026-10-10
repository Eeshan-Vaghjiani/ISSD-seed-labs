"""Build the Member 6 submission from verified local artifacts (HOST only).

Uses the sibling Member 5 deck primitives for the shared visual style. It never
boots a VM, runs lab programs, or reads host account/target files. Original
evidence copies refuse differing existing destinations; crops are documented.
"""
from pathlib import Path
import argparse
import hashlib
import importlib.util
import json
import re
import shutil
import subprocess
import tarfile

from PIL import Image
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pptx import Presentation
from pptx.dml.color import RGBColor as PColor
from pptx.util import Inches as PI


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'submission'
FRESH = ROOT / 'evidence/incoming/opus-fresh-20261008'
DOCUMENTS = ('REPORT', 'LIVE_DEMO', 'LIVE_COMMANDS', 'LIVE_PREPARATION', 'SLIDE_NOTES')
NAVY, TEAL, GRAY, PALE = '142A43', '087E8B', '526175', 'F2F5F8'

SELECTED = (
    'M6-S01-vm-setup.png', 'M6-S01-c-vm-setup.png', 'M6-S01-d-vm-setup.png',
    'M6-S02-environment.png', 'M6-S02-b-environment.png',
    'M6-S03-code-build.png', 'M6-S03-b-code-build.png', 'M6-S03-c-code-build.png',
    'M6-S04-dummy-baseline.png', 'M6-S05-normal-cow.png',
    'M6-S06-task1-running.png', 'M6-S07-task1-result.png',
    'M6-S08-charlie-baseline.png', 'M6-S08-b-charlie-baseline.png',
    'M6-S09-task2-uid-change.png', 'M6-S09-b-task2-uid-change.png',
    'M6-S10-task2-root-proof.png', 'M6-S10-b-proof-shell-exited.png',
    'M6-S11-account-restored.png', 'M6-S11-b-baseline-restored.png',
    'M6-S13-cleanup.png',
)

CROPS = {
    'report-S01-profile': (SELECTED[0], (212, 75, 910, 306)),
    'report-S01-cpu': (SELECTED[1], None),
    'report-S01-disk': (SELECTED[2], None),
    'report-S02': (SELECTED[3], (65, 25, 1276, 525)),
    'report-S02-tools': (SELECTED[4], (65, 25, 1276, 565)),
    'report-S03-build': (SELECTED[5], (65, 25, 1276, 545)),
    'report-S03-mapping': (SELECTED[6], (65, 25, 1276, 792)),
    'report-S03-workers': (SELECTED[7], (65, 25, 1276, 661)),
    'report-S04': (SELECTED[8], (65, 25, 1276, 375)),
    'report-S05': (SELECTED[9], (65, 25, 1276, 316)),
    'report-S06': (SELECTED[10], (65, 25, 1276, 506)),
    'report-S07': (SELECTED[11], (65, 25, 1276, 489)),
    'report-S08-login': (SELECTED[12], (65, 25, 1276, 375)),
    'report-S08-backup': (SELECTED[13], (65, 25, 1276, 432)),
    'report-S09-run': (SELECTED[14], (65, 25, 1276, 602)),
    'report-S09-diff': (SELECTED[15], (65, 25, 1276, 508)),
    'report-S10': (SELECTED[16], (65, 25, 1276, 414)),
    'report-S11': (SELECTED[18], (65, 25, 1276, 563)),
    'report-S13': (SELECTED[20], (65, 25, 1276, 489)),
    # Complete account/su/identity lines, omitting the earlier long seed groups.
    'slide-S10-proof': (SELECTED[16], (65, 105, 920, 413)),
}


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def write_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, indent=2, ensure_ascii=False) + '\n')


def copy_original(source, destination):
    destination.parent.mkdir(parents=True, exist_ok=True)
    if destination.exists():
        if sha(source) != sha(destination):
            raise RuntimeError('Preserve different existing artifact: ' + str(destination))
    else:
        shutil.copy2(source, destination)


def assets():
    selections = []
    for name in SELECTED:
        source = FRESH / name
        digest = sha(source)
        sidecar = source.with_suffix('.capture.json')
        if not name.startswith('M6-S01-'):
            capture = json.loads(sidecar.read_text())
            assert capture['sha256'] == digest, name
            assert capture['uuid'] == '242db196-83e1-4b7a-9f0b-e399d6f8b9dc'
            copy_original(sidecar, OUT / 'evidence' / sidecar.name)
        copy_original(source, OUT / 'evidence' / name)
        with Image.open(source) as image:
            dimensions = list(image.size)
            image.verify()
        selections.append({
            'selected': name, 'original': str(source.relative_to(ROOT)),
            'sha256': digest, 'dimensions': dimensions,
            'source_type': 'User-supplied host capture' if name.startswith('M6-S01-')
                           else 'Unedited VirtualBox guest framebuffer',
            'review': 'Visually inspected; see fresh EVIDENCE_REVIEW.md',
        })
    write_json(OUT / 'evidence/SELECTION.json', selections)

    manifest = []
    (OUT / 'figures').mkdir(parents=True, exist_ok=True)
    for name, (original, box) in CROPS.items():
        source = OUT / 'evidence' / original
        with Image.open(source) as image:
            rectangle = box or (0, 0, image.width, image.height)
            x0, y0, x1, y1 = rectangle
            assert 0 <= x0 < x1 <= image.width and 0 <= y0 < y1 <= image.height
            destination = OUT / 'figures' / (name + '.png')
            image.crop(rectangle).save(destination)
        manifest.append({
            'figure': destination.name, 'original': original,
            'original_sha256': sha(source), 'crop_box_xyxy': rectangle,
            'figure_sha256': sha(destination),
            'purpose': 'Labelled readability excerpt; contained pixels unchanged',
        })
    write_json(OUT / 'figures/CROP_MANIFEST.json', manifest)

    archive = FRESH / 'member6-fresh-final-20261009.tar.gz'
    assert sha(archive) == archive.with_name(archive.name + '.sha256').read_text().split()[0]
    with tarfile.open(archive, 'r:gz') as tar:
        for line in (ROOT / 'lab-files/SOURCE_SHA256SUMS').read_text().splitlines():
            digest, name = line.split()
            assert sha(ROOT / 'lab-files' / name) == digest
            assert hashlib.sha256(tar.extractfile('lab-files/' + name).read()).hexdigest() == digest
            copy_original(ROOT / 'lab-files' / name, OUT / 'source' / name)
        copy_original(ROOT / 'lab-files/SOURCE_SHA256SUMS', OUT / 'source/SOURCE_SHA256SUMS')
        for member in tar.getmembers():
            if not member.name.startswith('automation/') or not member.isfile():
                continue
            assert '/' not in member.name[len('automation/'):]
            data = tar.extractfile(member).read()
            destination = OUT / member.name
            destination.parent.mkdir(parents=True, exist_ok=True)
            if destination.exists():
                assert destination.read_bytes() == data, destination
            else:
                destination.write_bytes(data)
    for source in (FRESH / 'logs').iterdir():
        if source.is_file():
            copy_original(source, OUT / 'logs' / source.name)
    for source in (FRESH / 'provenance').iterdir():
        if source.is_file() and source.suffix in ('.json', '.jsonl', '.sha256', '.txt'):
            copy_original(source, OUT / 'provenance' / source.name)
    print(f'Assets: {len(selections)} original PNGs, {len(manifest)} documented crops; sources/logs verified.')


def shade(paragraph, fill):
    element = OxmlElement('w:shd')
    element.set(qn('w:fill'), fill)
    paragraph._p.get_or_add_pPr().append(element)


def inline(paragraph, text):
    text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1', text)
    for part in re.split(r'(\*\*.*?\*\*|`[^`]+`)', text):
        run = paragraph.add_run(part[2:-2] if part.startswith('**') else
                                part[1:-1] if part.startswith('`') else part)
        if part.startswith('**'):
            run.bold = True
        if part.startswith('`'):
            run.font.name = 'Liberation Mono'
            run.font.size = Pt(9)


def page_field(paragraph, instruction):
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), instruction)
    field.set(qn('w:dirty'), 'true')
    run = OxmlElement('w:r')
    text = OxmlElement('w:t')
    text.text = '1'
    run.append(text)
    field.append(run)
    paragraph._p.append(field)


def render_doc(name):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.2677), Inches(11.6929)
    section.top_margin = section.bottom_margin = Inches(.62)
    section.left_margin = section.right_margin = Inches(.65)
    section.header_distance = section.footer_distance = Inches(.27)
    compact = name == 'LIVE_COMMANDS'
    for style_name in ('Normal', 'Body Text', 'List Bullet', 'List Number'):
        style = doc.styles[style_name]
        style.font.name = 'Liberation Sans'
        style.font.size = Pt(9.5 if compact else 10 if name == 'LIVE_DEMO' else 10.5)
        style.paragraph_format.space_after = Pt(3 if compact else 4 if name == 'LIVE_DEMO' else 6)
        style.paragraph_format.line_spacing = 1 if compact else 1.06
    for level, size in ((1, 17), (2, 13), (3, 11)):
        style = doc.styles['Heading ' + str(level)]
        style.font.name = 'Liberation Sans'
        style.font.size = Pt(min(size, 13) if compact else size)
        style.font.color.rgb = RGBColor.from_string(NAVY if level == 1 else TEAL)
        style.paragraph_format.space_before = Pt(7 if compact else 12)
        style.paragraph_format.space_after = Pt(4 if compact else 6)
        style.paragraph_format.keep_with_next = True
    doc.styles['Title'].font.name = 'Liberation Sans'
    doc.styles['Title'].font.size = Pt(24 if compact else 25 if name == 'LIVE_PREPARATION' else 29)
    doc.styles['Title'].font.color.rgb = RGBColor.from_string(NAVY)
    header = section.header.paragraphs[0]
    header.text = 'ISSD  /  MEMBER 6                                         SEED DIRTY COW LAB'
    for run in header.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(GRAY)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Member 6 — ISSD   |   ')
    page_field(footer, 'PAGE')
    footer.add_run(' / ')
    page_field(footer, 'NUMPAGES')
    for run in footer.runs:
        run.font.size = Pt(8)
    doc.core_properties.author = 'Member 6 — ISSD'
    doc.core_properties.title = name.replace('_', ' ').title()
    doc.core_properties.subject = 'Verified Dirty COW findings and live demonstration support'

    lines = (OUT / (name + '.md')).read_text().splitlines()
    index = 0
    while index < len(lines):
        line = lines[index]
        if line == '<!-- pagebreak -->':
            doc.add_page_break()
        elif line.startswith('```'):
            code = []
            index += 1
            while index < len(lines) and not lines[index].startswith('```'):
                code.append(lines[index])
                index += 1
            assert index < len(lines), 'Unclosed code fence: ' + name
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.left_indent = Inches(.09)
            paragraph.paragraph_format.right_indent = Inches(.06)
            paragraph.paragraph_format.space_before = Pt(2 if compact else 4)
            paragraph.paragraph_format.space_after = Pt(4 if compact else 8)
            paragraph.paragraph_format.line_spacing = 1
            shade(paragraph, 'EEF2F6')
            run = paragraph.add_run('\n'.join(code))
            run.font.name = 'Liberation Mono'
            run.font.size = Pt(8.4 if name == 'REPORT' else 9)
        elif match := re.match(r'^!\[(.*?)\]\((.*?)\)$', line):
            caption, filename = match.groups()
            image_path = OUT / filename
            with Image.open(image_path) as image:
                width = min(6.90, 7.8 * image.width / image.height)
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True
            picture = paragraph.add_run().add_picture(str(image_path), width=Inches(width))
            picture._inline.docPr.set('descr', caption)
            paragraph = doc.add_paragraph(caption)
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.space_after = Pt(12)
            for run in paragraph.runs:
                run.italic = True
                run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor.from_string(GRAY)
        elif line.startswith('|'):
            rows = []
            while index < len(lines) and lines[index].startswith('|'):
                row = [cell.strip() for cell in re.split(
                    r'\|(?=(?:[^`]*`[^`]*`)*[^`]*$)', lines[index].strip().strip('|'))]
                if not all(re.fullmatch(r':?-+:?', cell) for cell in row):
                    rows.append(row)
                index += 1
            table = doc.add_table(rows=0, cols=max(map(len, rows)))
            table.style = 'Light Shading Accent 1'
            for row_index, row in enumerate(rows):
                cells = table.add_row().cells
                for col, value in enumerate(row):
                    paragraph = cells[col].paragraphs[0]
                    paragraph.paragraph_format.space_before = Pt(3)
                    paragraph.paragraph_format.space_after = Pt(3)
                    inline(paragraph, value)
                    for run in paragraph.runs:
                        run.font.size = Pt(9)
                        if row_index == 0:
                            run.bold = True
                properties = table.rows[row_index]._tr.get_or_add_trPr()
                properties.append(OxmlElement('w:cantSplit'))
                if row_index == 0:
                    properties.append(OxmlElement('w:tblHeader'))
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
            continue
        elif match := re.match(r'^(#{1,4}) (.*)$', line):
            depth, text = len(match[1]), match[2]
            paragraph = doc.add_paragraph(style='Title') if depth == 1 else doc.add_heading(level=depth - 1)
            inline(paragraph, text)
        elif line.startswith('* '):
            inline(doc.add_paragraph(style='List Bullet'), line[2:])
        elif line.strip() and not line.startswith('<!--'):
            paragraph = doc.add_paragraph()
            if line.startswith('**') and line.endswith('**') or line.rstrip().endswith(':'):
                paragraph.paragraph_format.keep_with_next = True
            inline(paragraph, line)
        index += 1
    target = OUT / (name + '.docx')
    doc.save(target)
    reopened = Document(target)
    assert len(reopened.inline_shapes) == len(doc.inline_shapes)
    print(target.name, len(reopened.paragraphs), 'paragraphs;', len(reopened.inline_shapes), 'images')


def build_deck():
    specification = importlib.util.spec_from_file_location(
        'member5_deck_primitives', ROOT.parent / 'Member5/tools/build_submission.py')
    shared = importlib.util.module_from_spec(specification)
    specification.loader.exec_module(shared)

    class Deck(shared.Deck):
        def text(self, *args, **kwargs):
            kwargs.setdefault('font', 'Liberation Sans')
            return super().text(*args, **kwargs)

        def base(self, title, subtitle='', tag='FINDINGS', backup=False):
            slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = PColor.from_string('FAFBFD')
            self.rect(slide, 0, 0, 13.3333, .12, TEAL, False)
            self.text(slide, .55, .31, 12, .25,
                      'ISSD  /  MEMBER 6  /  ' + ('BACKUP — ' if backup else '') + tag,
                      10, TEAL, True)
            self.text(slide, .55, .76, 12.15, .65, title, 32, NAVY, True)
            if subtitle:
                self.text(slide, .57, 1.48, 12.0, .57, subtitle, 17, GRAY)
            self.text(slide, .56, 7.07, 11.5, .19,
                      'SEED Dirty COW • Wenliang Du / SEED Labs • Recorded findings: 9 Oct 2026', 9, GRAY)
            self.text(slide, 12.20, 7.02, .55, .29, str(len(self.prs.slides)), 12, TEAL, True)
            return slide

    notes = {}
    for match in re.finditer(r'^## Slide (\d+) — (.*?)\n(.*?)(?=^## |\Z)',
                             (OUT / 'SLIDE_NOTES.md').read_text(), re.M | re.S):
        notes[int(match[1])] = (match[2], match[3].strip())
    deck = Deck()
    slide = deck.prs.slides.add_slide(deck.prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = PColor.from_string(NAVY)
    deck.rect(slide, .64, .64, .14, 6.00, TEAL, False)
    deck.text(slide, 1.04, .72, 11.0, .38, 'MEMBER 6 — ISSD', 16, '79DDD7', True)
    deck.text(slide, 1.03, 1.61, 11.35, 1.65, 'Private should\nstay private.', 48, 'FFFFFF', True)
    deck.text(slide, 1.06, 3.70, 10.9, .53, 'Dirty COW • CVE-2016-5195', 29, 'FFFFFF')
    deck.text(slide, 1.07, 4.53, 10.65, .84,
              'A local kernel race that changed protected file bytes\nand enabled a fresh UID-0 login.', 24, 'DDE9F4')
    deck.rect(slide, 1.04, 5.94, 11.18, .61, '203E5E')
    deck.text(slide, 1.29, 6.08, 10.64, .34,
              'COW  →  LIVE DUMMY DEMO  →  RECORDED ACCOUNT PROOF  →  DEFENCES', 16, '8DE4DF', True)

    slide = deck.base('Where private isolation breaks',
                      'The control and the exploit use different memory-write paths.', 'MECHANISM')
    deck.rect(slide, .61, 2.20, 5.94, 3.77)
    deck.text(slide, .90, 2.45, 5.37, .45, 'Normal COW control', 27, TEAL, True)
    deck.text(slide, .91, 3.10, 5.24, 1.56,
              'O_RDONLY file\nWritable MAP_PRIVATE mapping\nmemcpy changes a private copy', 22)
    deck.text(slide, .91, 5.13, 5.27, .51, 'Backing file unchanged', 24, TEAL, True)
    deck.rect(slide, 6.80, 2.20, 5.93, 3.77, NAVY)
    deck.text(slide, 7.09, 2.45, 5.38, .45, 'Dirty COW race', 27, '8DE4DF', True)
    deck.text(slide, 7.10, 3.10, 5.24, 1.59,
              'O_RDONLY + read-only MAP_PRIVATE\nWriter: pwrite to /proc/self/mem\nDiscard: MADV_DONTNEED', 21, 'FFFFFF')
    deck.text(slide, 7.10, 5.13, 5.26, .55, 'Vulnerable kernel → file changes', 22, 'F6CE84', True)
    deck.text(slide, .82, 6.33, 11.8, .48,
              'pwrite uses a virtual process address. The race is inside kernel COW/page handling.', 18, GRAY)

    slide = deck.base('Normal file permissions were already restrictive',
                      'Recorded environment: Ubuntu 12.04.2 • i686 / 32-bit • kernel 3.5.0-37-generic', 'ENVIRONMENT')
    deck.rect(slide, .63, 2.13, 12.04, .65, NAVY)
    deck.text(slide, .86, 2.29, 11.56, .37,
              'seed UID 1000   |   root:root 0644 target   |   seed-owned 0755 binaries   |   native ext4',
              19, 'FFFFFF', True)
    deck.picture(slide, OUT / 'figures/report-S04.png', .76, 3.03, 11.81, 3.47)
    deck.text(slide, .83, 6.61, 11.65, .25,
              'Actual S04 excerpt. Package: 3.5.0-37.58~precise1 / linux-lts-quantal. No Set-UID installation.', 13, GRAY)

    slide = deck.base('Switch to the live dummy demonstration',
                      'Prepare Part 0 first. Use the same guest terminal and a new run label.', 'LIVE DEMO')
    for index, (title, body) in enumerate((
        ('1  Protection', 'id • permissions • original file\nOrdinary write should be denied.'),
        ('2  Normal COW', './cow_control\nPrivate stars; backing file original.'),
        ('3  Bounded race', '15-second classroom trial\nRead and verify the whole file.'),
    )):
        x = .64 + index * 4.17
        deck.rect(slide, x, 2.27, 3.91, 2.69)
        deck.text(slide, x + .24, 2.54, 3.43, .50, title, 25, TEAL, True)
        deck.text(slide, x + .24, 3.30, 3.39, 1.24, body, 21)
    deck.rect(slide, .68, 5.39, 11.99, .66, NAVY)
    deck.text(slide, .96, 5.56, 11.37, .34,
              'bash run_trial.sh dummy 15 "$RUN"', 24, 'FFFFFF', font='Liberation Mono')
    deck.text(slide, .84, 6.35, 11.72, .42,
              'Timing miss? Preserve it and show recorded S07: exact success, 0 integer seconds / 30-second bound.',
              16, GRAY)

    slide = deck.base('Recorded file change → fresh UID-0 login',
                      'Task 2: only charlie’s UID field changed. All other bytes and root:root 0644 were retained.', 'ACCOUNT IMPACT')
    deck.rect(slide, .66, 2.15, 12.00, .57, NAVY)
    deck.text(slide, .95, 2.28, 11.37, .34,
              'Original UID 1001   →   field 0000   →   fresh non-sudo login: UID 0   |   GID stays 1002',
              19, 'FFFFFF', True)
    deck.picture(slide, OUT / 'figures/slide-S10-proof.png', 1.04, 2.97, 11.22, 3.39)
    deck.text(slide, .84, 6.57, 11.72, .29,
              'S10 excerpt, 9 Oct 2026. First authentication failed; retry succeeded. Full original and recordings retained.',
              13, GRAY)

    slide = deck.base('Fix the kernel; verify the restored state',
                      'All four required countermeasure themes. Optional patched comparison: not performed; discussion only.', 'DEFENCES')
    cards = (
        ('Patch the kernel', 'Install the vendor fix and confirm\nthe fixed kernel is actually booted.'),
        ('Maintain updates', 'Use a supported OS and track\npackage-family backports.'),
        ('Limit local execution', 'Reduce unnecessary local access\nand untrusted code execution.'),
        ('Monitor escalation', 'Check account integrity, numeric\nUID-0 entries and privileged sessions.'),
    )
    for index, (title, body) in enumerate(cards):
        x, y = .67 + (index % 2) * 6.24, 2.24 + (index // 2) * 1.72
        deck.rect(slide, x, y, 5.76, 1.56)
        deck.text(slide, x + .24, y + .18, 5.20, .42, title, 24, TEAL, True)
        deck.text(slide, x + .24, y + .66, 5.21, .74, body, 19)
    deck.rect(slide, .69, 5.94, 11.96, .76, NAVY)
    deck.text(slide, .93, 6.13, 11.41, .40,
              'Verified afterward: old root shell ended • exact baseline • new UID 1001 login • cleanup',
              18, 'FFFFFF', True)

    slide = deck.base('Recorded Task 1: the actual backing file changed',
                      'fresh-t1-01 • 30-second limit • 0 integer elapsed seconds • no finer runtime or attempt count measured',
                      'DUMMY RESULT', True)
    deck.picture(slide, OUT / 'figures/report-S07.png', .79, 2.20, 11.76, 4.35)
    deck.text(slide, .86, 6.62, 11.60, .27,
              'S07 readability excerpt: complete diff, fresh read, root:root 0644 / 19 bytes and exact-byte comparison.', 13, GRAY)

    slide = deck.base('Recorded Task 2: exact UID-only replacement',
                      'One ordinary-seed trial. 2040 bytes before and after; four-character UID field at byte offset 2002.',
                      'ACCOUNT DIFF', True)
    deck.picture(slide, OUT / 'figures/report-S09-diff.png', .79, 2.17, 11.77, 4.42)
    deck.text(slide, .86, 6.63, 11.60, .26,
              'S09b: live file equals the saved after-copy; full-file verification separately confirms no other changes.', 13, GRAY)

    slide = deck.base('Restoration is a fresh-login check too',
                      'Proof shell PID 3392 was exited first. Exact normal bytes and a new ordinary login both passed.',
                      'RESTORATION', True)
    deck.picture(slide, OUT / 'figures/report-S11.png', .86, 2.14, 11.60, 4.43)
    deck.text(slide, .87, 6.64, 11.55, .25,
              'S11 excerpt. Existing process credentials survive file edits; close privileged shells before restoration.', 13, GRAY)

    slide = deck.base('Sources, evidence and assistance',
                      'Completed practical evidence is distinct from the future classroom demonstration.', 'REFERENCES', True)
    deck.text(slide, .85, 2.26, 11.66, 2.45,
              'Wenliang Du / SEED Labs — Dirty COW Attack Lab (2017)\n'
              'seedsecuritylabs.org/Labs_20.04/Software/Dirty_COW/\n'
              'Historical SEED12 manual • mmap(2) • madvise(2) • proc_pid_mem(5)\n'
              'Ubuntu CVE-2016-5195 information • supplied Member 6 allocation', 22, NAVY)
    deck.rect(slide, .80, 5.02, 11.79, .81)
    deck.text(slide, 1.04, 5.23, 11.20, .36,
              'REPORT.pdf   |   LIVE_DEMO.pdf   |   SLIDE_NOTES.pdf   |   video/TASK2_FALLBACK.mp4', 18, TEAL, True)
    deck.text(slide, .87, 6.11, 11.57, .62,
              'CC BY-NC-SA 4.0 attribution retained. GPT-6 Astra / OpenCode assisted execution and preparation;\n'
              'authentication was interactive. No oral rehearsal or course upload is claimed.', 15, GRAY)

    assert len(deck.prs.slides) == len(notes) == 10
    for number, slide in enumerate(deck.prs.slides, 1):
        heading, text = notes[number]
        deck.note(slide, heading, text)
    for item in deck.bounds:
        x, y, width, height = item['box']
        assert x >= 0 and y >= 0 and x + width <= 13.34 and y + height <= 7.51, item
    deck.prs.core_properties.title = 'Dirty COW — findings and live demonstration'
    deck.prs.core_properties.author = 'Member 6 — ISSD'
    deck.prs.core_properties.subject = 'Six main slides plus four evidence/reference backups; recorded findings 9 October 2026'
    target = OUT / 'MAIN_PRESENTATION.pptx'
    deck.prs.save(target)
    assert len(Presentation(target).slides) == 10
    write_json(OUT / 'provenance/SLIDE_LAYOUT.json', deck.bounds)
    print(target.name, '10 slides with embedded notes')


def export_pdf():
    profile = Path('/tmp/opencode/member6-libreoffice-profile')
    command = [
        'libreoffice', '-env:UserInstallation=' + profile.as_uri(), '--headless',
        '--convert-to', 'pdf', '--outdir', str(OUT),
        *[str(OUT / (name + '.docx')) for name in DOCUMENTS],
        str(OUT / 'MAIN_PRESENTATION.pptx'),
    ]
    result = subprocess.run(command, text=True, capture_output=True)
    write_json(OUT / 'provenance/PDF_EXPORT.json', {
        'command': command, 'returncode': result.returncode,
        'stdout': result.stdout, 'stderr': result.stderr,
    })
    print(result.stdout)
    result.check_returncode()
    for name in (*DOCUMENTS, 'MAIN_PRESENTATION'):
        assert (OUT / (name + '.pdf')).is_file(), name


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--only', choices=('all', 'assets', 'docs', 'slides', 'pdf'), default='all')
    args = parser.parse_args()
    if args.only in ('all', 'assets'):
        assets()
    if args.only in ('all', 'slides'):
        build_deck()
    if args.only in ('all', 'docs'):
        for name in DOCUMENTS:
            render_doc(name)
    if args.only in ('all', 'pdf'):
        export_pdf()


if __name__ == '__main__':
    main()
