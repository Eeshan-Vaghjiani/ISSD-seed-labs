"""HOST-only integrity, export and layout audit of the finished Member 6 pack.

Reads submitted copies and fresh source artifacts; never runs a lab program or
uses host /zzz, /etc/passwd, or account commands. Visual review is separate.
"""
from pathlib import Path
import argparse
import ast
import datetime as dt
import hashlib
import json
import re
import subprocess
import zipfile

from PIL import Image, ImageChops, ImageDraw
from docx import Document
from pptx import Presentation
import pymupdf as fitz


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'submission'
FRESH = ROOT / 'evidence/incoming/opus-fresh-20261008'
DOCUMENTS = ('REPORT', 'LIVE_DEMO', 'LIVE_COMMANDS', 'LIVE_PREPARATION', 'SLIDE_NOTES')


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def normalize(text):
    return re.sub(r'\s+', '', text.replace('\u00ad', '').replace('\u200b', '').replace('`', ''))


def render_review(directory):
    directory.mkdir(parents=True, exist_ok=True)
    for name in (*DOCUMENTS, 'MAIN_PRESENTATION'):
        document = fitz.open(OUT / (name + '.pdf'))
        thumbnails = []
        for index, page in enumerate(document):
            pixmap = page.get_pixmap(matrix=fitz.Matrix(1.4, 1.4), alpha=False)
            path = directory / f'{name}-{index + 1:02}.png'
            pixmap.save(path)
            with Image.open(path) as image:
                thumbnail = image.convert('RGB')
                thumbnail.thumbnail((440, 580) if name != 'MAIN_PRESENTATION' else (620, 350))
                thumbnails.append((index + 1, thumbnail))
        page_width, page_height = (470, 620) if name != 'MAIN_PRESENTATION' else (650, 385)
        for start in range(0, len(thumbnails), 6):
            subset = thumbnails[start:start + 6]
            rows = (len(subset) + 1) // 2
            sheet = Image.new('RGB', (page_width * 2, page_height * rows), '#D9E1E9')
            draw = ImageDraw.Draw(sheet)
            for slot, (number, thumbnail) in enumerate(subset):
                x, y = (slot % 2) * page_width, (slot // 2) * page_height
                draw.text((x + 14, y + 8), f'{name} — page {number}', fill='black')
                sheet.paste(thumbnail, (x + (page_width - thumbnail.width) // 2, y + 30))
            sheet.save(directory / f'{name}-contact-{start // 6 + 1:02}.png')
        print('Rendered', name, len(document), 'pages for visual review.')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--render', action='store_true')
    parser.add_argument('--record', action='store_true',
                        help='Refresh provenance/SUBMISSION_VERIFICATION.json after a build')
    args = parser.parse_args()
    findings = []

    selected = json.loads((OUT / 'evidence/SELECTION.json').read_text())
    assert len(selected) == 21
    for item in selected:
        source = ROOT / item['original']
        copy = OUT / 'evidence' / item['selected']
        assert sha(source) == sha(copy) == item['sha256'], copy
        with Image.open(copy) as image:
            image.verify()
    findings.append('21 selected original PNGs match fresh sources and recorded hashes.')

    crops = json.loads((OUT / 'figures/CROP_MANIFEST.json').read_text())
    for item in crops:
        original, crop = OUT / 'evidence' / item['original'], OUT / 'figures' / item['figure']
        assert sha(original) == item['original_sha256'] and sha(crop) == item['figure_sha256']
        with Image.open(original) as image, Image.open(crop) as actual:
            expected = image.crop(item['crop_box_xyxy']).convert('RGB')
            assert expected.size == actual.size
            assert ImageChops.difference(expected, actual.convert('RGB')).getbbox() is None
    findings.append(f'{len(crops)} labelled crops preserve the exact pixels of their documented rectangles.')

    for line in (OUT / 'source/SOURCE_SHA256SUMS').read_text().splitlines():
        digest, name = line.split()
        assert sha(OUT / 'source' / name) == sha(ROOT / 'lab-files' / name) == digest
    for path in (OUT / 'logs').iterdir():
        if path.is_file():
            assert sha(path) == sha(FRESH / 'logs' / path.name), path
    logs = OUT / 'logs'
    for label, metadata in (('fresh-t1-01', b'0:0:644:19\n'), ('fresh-t2-01', b'0:0:644:2040\n')):
        assert (logs / (label + '-metadata-before.txt')).read_bytes() == metadata
        assert (logs / (label + '-metadata-after.txt')).read_bytes() == metadata
        text = (logs / (label + '-verification.log')).read_text()
        result = json.loads(text[text.index('{'):])
        assert result['ACTUAL'] == 'exact-change' and result['live_file_checked']
        assert result['ruid'] == result['euid'] == 1000
        assert result['elapsed_integer_seconds'] == 0 and result['kernel_race_attempt_count'] is None
    assert (logs / 'fresh-t1-01-before.txt').read_bytes() == b'111111222222333333\n'
    assert (logs / 'fresh-t1-01-after.txt').read_bytes() == b'111111******333333\n'
    before = (logs / 'fresh-t2-01-before.txt').read_bytes()
    after = (logs / 'fresh-t2-01-after.txt').read_bytes()
    original_record = b'charlie:x:1001:1002:,,,:/home/charlie:/bin/bash\n'
    assert before.count(original_record) == 1
    assert after == before.replace(original_record, original_record.replace(b':1001:', b':0000:', 1), 1)
    assert (logs / 'fresh-restoration-01-passwd.txt').read_bytes() == before
    assert hashlib.sha256(before).hexdigest() == '3edf14347c28c11d3a8816326e14baa3b99b569f1dc33dd15cd7d953c91dc337'
    transcript = (logs / 'fresh-task2-20261009.typescript').read_bytes()
    assert b'su: Authentication failure' in transcript and b'Proof shell PID=3392' in transcript
    assert b'uid=0(root) gid=1002(charlie)' in transcript
    assert b'uid=1001(charlie) gid=1002(charlie)' in transcript
    assert b'Checked: exact normal passwd/charlie/UID-0 baseline, no attacker/su/root shell, no /zzz.' in (logs / 'fresh-cleanup-01.txt').read_bytes()
    findings.append('Original sources, logs, full trial changes, restored file, identity chain and cleanup artifacts are consistent.')

    for path in (OUT / 'automation').glob('*.py'):
        ast.parse(path.read_text(), filename=str(path))
    for path in (OUT / 'automation').glob('*.sh'):
        subprocess.run(['bash', '-n', str(path)], check=True)

    document_records = []
    for name in DOCUMENTS:
        path = OUT / (name + '.docx')
        with zipfile.ZipFile(path) as archive:
            assert archive.testzip() is None
        document = Document(path)
        pdf = fitz.open(OUT / (name + '.pdf'))
        # Font fallback can put an arrow in a separate PDF text object. Check
        # both stream order and geometric reading order, retaining all glyphs.
        bodies = [normalize(''.join(page.get_text(
            clip=fitz.Rect(0, 35, page.rect.width, page.rect.height - 35), sort=sort)
            for page in pdf)) for sort in (False, True)]
        missing = [paragraph.text for paragraph in document.paragraphs
                   if paragraph.text.strip() and not any(normalize(paragraph.text) in body for body in bodies)]
        missing += [cell.text for table in document.tables for row in table.rows for cell in row.cells
                    if cell.text.strip() and not any(normalize(cell.text) in body for body in bodies)]
        assert not missing, f'{name}: missing exported text: {missing[:4]}'
        markdown = (OUT / (name + '.md')).read_text()
        assert not re.search(r'\[(?:INSERT|OBSERVED|PASS/FAIL)\b', markdown), name
        assert len(document.inline_shapes) == len(re.findall(r'^!\[', markdown, re.M))
        for target in re.findall(r'^!\[[^\]]*\]\(([^)]+)\)', markdown, re.M):
            assert (OUT / target).is_file(), target
        for page in pdf:
            for word in page.get_text('words'):
                assert page.rect.contains(fitz.Rect(word[:4])), f'{name}: off-page word on page {page.number + 1}: {word}'
        document_records.append({'name': name, 'pages': len(pdf), 'images': len(document.inline_shapes),
                                 'docx_sha256': sha(path), 'pdf_sha256': sha(OUT / (name + '.pdf'))})
        findings.append(f'{name}: {len(pdf)} PDF pages; {len(document.inline_shapes)} images; matching DOCX/PDF text; no off-page words/placeholders.')

    deck = Presentation(OUT / 'MAIN_PRESENTATION.pptx')
    pdf = fitz.open(OUT / 'MAIN_PRESENTATION.pdf')
    assert len(deck.slides) == len(pdf) == 10
    for number, slide in enumerate(deck.slides):
        texts = [normalize(pdf[number].get_text(sort=sort)) for sort in (False, True)]
        assert slide.notes_slide.notes_text_frame.text.strip()
        for shape in slide.shapes:
            if shape.has_text_frame:
                assert any(normalize(shape.text) in text for text in texts), f'Slide {number + 1} text mismatch: {shape.text}'
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= deck.slide_width + 1000
            assert shape.top + shape.height <= deck.slide_height + 1000
        for word in pdf[number].get_text('words'):
            assert pdf[number].rect.contains(fitz.Rect(word[:4])), f'Slide {number + 1} off-page word'
    findings.append('10-slide PPTX/PDF text matches; every slide has embedded notes; all shapes/words within page bounds.')

    video = OUT / 'video'
    manifest = json.loads((video / 'VIDEO_MANIFEST.json').read_text())
    assert sha(video / manifest['output']) == manifest['output_sha256']
    assert manifest['full_decode_passed']
    for item in manifest['originals']:
        assert sha(video / 'originals' / item['filename']) == sha(FRESH / item['filename']) == item['sha256']
    assert manifest['ranges'] == [
        ['M6-task2-recorded-screen0.webm', 0.0, 22.0, '01-uid-overwrite.mp4'],
        ['M6-task2-login-screen0.webm', 268.0, 304.2, '02-login-proof.mp4'],
    ]
    findings.append('Video original hashes, explicit edit intervals and fully decoded H.264 fallback verified.')

    record = {
        'verified_utc': dt.datetime.now(dt.timezone.utc).isoformat(),
        'boundary': 'HOST: local artifact/package checks only',
        'findings': findings, 'documents': document_records,
        'presentation_sha256': sha(OUT / 'MAIN_PRESENTATION.pptx'),
        'presentation_pdf_sha256': sha(OUT / 'MAIN_PRESENTATION.pdf'),
        'video_sha256': manifest['output_sha256'],
        'visual_review': 'Separate manual inspection of rendered pages/slides is documented in BUILD_VERIFICATION.md.',
    }
    if args.record:
        (OUT / 'provenance/SUBMISSION_VERIFICATION.json').write_text(json.dumps(record, indent=2) + '\n')
    for finding in findings:
        print('PASS:', finding)
    if args.render:
        render_review(Path('/tmp/opencode/member6-layout-review'))


if __name__ == '__main__':
    main()
